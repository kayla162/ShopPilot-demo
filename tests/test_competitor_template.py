import io
from datetime import date

import openpyxl
import pandas as pd
import pytest

from shoppilot import db, pipeline
from shoppilot.analytics import competitor_template as ct
from shoppilot.ingest import detector, excel_reader


def _products():
    return pd.DataFrame({
        "product_id": ["1", "2", "3"], "code": ["A9991", "A9994", None],
        "core_title": ["細肩吊帶排釦背心示範(附腰帶)", "格紋細肩吊帶排釦背心示範(附腰帶)", "直播自助下單區"],
        "search_keyword": ["細肩 吊帶 排釦 背心", "格紋 細肩 吊帶 背心", "直播自助下單區"],
        "search_url": ["https://shopee.tw/search?keyword=a", "https://shopee.tw/search?keyword=b", "u"],
        "is_special": [False, False, True], "has_price": [True, True, True], "impressions": [5688, 710, 50]})


def _filled(rows) -> bytes:
    """模擬賣家填好的範本。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ct.SHEET
    ws.append(ct.TEMPLATE_COLUMNS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_file_name_and_default_ids():
    assert ct.template_file_name(date(2026, 9, 30)) == "shoppilot_競品範本_20260930.xlsx"
    assert ct.default_product_ids(_products(), top_n=1) == ["1"]
    assert ct.default_product_ids(_products()) == ["1", "2"]           # 特殊商品不列入


def test_build_template_prefills_keywords_and_links():
    data = ct.build_template(_products(), ["1", "2"])
    rows = excel_reader.rows(io.BytesIO(data), ct.SHEET)
    assert rows[0] == ct.TEMPLATE_COLUMNS
    assert rows[1][:4] == ["1", "A9991 細肩吊帶排釦背心示範(附腰帶)", "細肩 吊帶 排釦 背心", "https://shopee.tw/search?keyword=a"]
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb[ct.SHEET].cell(2, 4).hyperlink.target == "https://shopee.tw/search?keyword=a"
    assert ct.HELP_SHEET in wb.sheetnames


def test_parse_template_forward_fills_and_reports_errors():
    data = _filled([
        ["1", "A9991 細肌…", "細肩 吊帶 排釦 背心", "u1"],
        [None, None, None, None, "A店 同款排釦馬甲", "$350", "已售出 1.2萬", "A店", "https://x"],
        [None, None, None, None, "B店 馬甲", "$250 - $399", "234", None, None, "否"],
        [None, None, None, None, "C店 馬甲", "問賣家"],
        ["2", None, None, None, "D店 格紋馬甲", 299, None, None, None, "是"],
    ])
    df, errors = ct.parse_template(io.BytesIO(data))
    assert df["my_product_id"].tolist() == ["1", "1", "2"]
    first = df.iloc[0]
    assert (first["price_min"], first["price_max"], first["sold_count"], first["search_keyword"]) == \
        (350.0, 350.0, 12000, "細肩 吊帶 排釦 背心")
    assert (df.iloc[1]["price_min"], df.iloc[1]["price_max"], df.iloc[1]["include_override"]) == (250.0, 399.0, "no")
    assert (df.iloc[2]["price_min"], df.iloc[2]["include_override"]) == (299.0, "yes")
    assert len(errors) == 1 and "問賣家" in errors[0]


def test_parse_template_rejects_wrong_files():
    buf = io.BytesIO()
    openpyxl.Workbook().save(buf)
    with pytest.raises(ValueError, match="競品清單"):
        ct.parse_template(io.BytesIO(buf.getvalue()))
    wb = openpyxl.Workbook()
    wb.active.title = ct.SHEET
    wb.active.append(["我的商品ID", "價格"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="競品名稱"):
        ct.parse_template(io.BytesIO(buf.getvalue()))


def test_competitor_upload_replaces_per_product():
    conn = db.connect(":memory:")
    try:
        first = _filled([["1", None, None, None, "A店", "$300"], [None, None, None, None, "B店", "$320"],
                         ["2", None, None, None, "C店", "$100"]])
        assert pipeline.import_file(conn, "我的競品.xlsx", first, force_type=detector.COMPETITOR).status == "imported"
        second = _filled([["1", None, None, None, "新的競品", "$310"]])
        assert pipeline.import_file(conn, "我的競品2.xlsx", second, force_type=detector.COMPETITOR).status == "imported"
        assert sorted(db.read_competitors(conn)["comp_name"]) == ["C店", "新的競品"]
    finally:
        conn.close()


def test_competitor_upload_without_rows_fails():
    conn = db.connect(":memory:")
    try:
        result = pipeline.import_file(conn, "空的.xlsx", _filled([["1"]]), force_type=detector.COMPETITOR)
        assert result.status == "failed" and "沒有可匯入的競品" in result.message
    finally:
        conn.close()


def test_help_sheet_states_inclusion_rule():
    assert any("留白或填「是」都會納入比較" in line and "填「否」" in line for line in ct.HELP_LINES)
