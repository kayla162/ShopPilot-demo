"""虛構賣場的示範報表：輸出和蝦皮相同格式的 O／P／B′ 報表與外部競品範本，給線上 demo 和測試用。

所有內容都是用固定亂數種子產生的虛構資料，和任何真實賣場無關：商品ID用 9 開頭、規格ID用 8 開頭、
主商品貨號用 D 開頭。標題裡不放貨號（transform/title.py 只認 A 開頭的貨號），畫面上用商品名稱標示。
數字刻意埋入幾個規律（標記多→曝光低、缺貨多→銷量低、同類偏貴→轉換低），讓曝光因子和規則建議有結果可以展示。
"""
from __future__ import annotations

import argparse
import math
import sys
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import openpyxl

from shoppilot import config, pipeline
from shoppilot.analytics import competitor_template as ct
from shoppilot.ingest import parsers
from shoppilot.transform.title import core_title

SHOP_NO = "900000001"
DEFAULT_TODAY = date(2026, 10, 1)
N_REGULAR = 150
DELETED_NAME = "由於商品被刪除而無法獲取訊息"

# 類別 → (商品數, 名詞, 基本價格)。名詞要能被 config/keywords.yaml 的 categories 分到這個類別
CATEGORIES = {
    "上衣": (34, ["上衣", "襯衫", "T恤", "針織衣"], 320), "外套": (12, ["外套", "開衫"], 590),
    "背心": (12, ["背心"], 250), "洋裝": (22, ["洋裝", "連身裙"], 520), "長裙": (10, ["長裙"], 450),
    "短裙": (12, ["短裙", "百褶裙"], 380), "長褲": (18, ["長褲", "寬褲"], 480), "短褲": (10, ["短褲"], 330),
    "套裝": (12, ["套裝", "兩件套"], 690), "配件": (8, ["髮飾", "帽"], 160),
}
SPECIALS = [("零碼清倉 福袋合集", 199), ("出清 特價區 單件價", 150), ("直播自助下單區", 150)]
MARKERS = ["【新品】", "【熱銷】", "#韓系穿搭", "#夏日必備", "🖤", "衝好評", "數量有限"]
MARKETING = ["激推!", "實拍"]
COLORS = ["黑色", "白色", "杏色", "粉色", "藍色", "灰色"]
SIZES = ["S", "M", "L"]
PREORDER_BASES = [100, 300, 500]
MODEL_INFO = "模特兒身高 165 公分、體重 48 公斤，穿 M 號。"
NOTICE = "🎉每樣商品都是實拍，顏色以實品為準。\n📌下單須知📌\n1. 下單前請先確認尺寸。\n2. 預購商品約 7–14 天出貨。"


@dataclass
class DemoSku:
    sku_id: str
    name: str
    sku_code: str
    price: int
    stock: int


@dataclass
class DemoProduct:
    pid: str
    code: str
    title: str
    description: str
    category: str
    special: bool
    preorder: bool
    overhead: float                      # 標記字數占比（產生器自己估的，用來決定曝光）
    skus: list[DemoSku] = field(default_factory=list)


def _attribute_pool(kw: dict) -> dict[str, list[str]]:
    """屬性詞去掉含有類別關鍵字的詞（例如「闊腿」「毛衣」），不然商品會被分到別的類別。"""
    cat_words = [w for c in kw["categories"] for w in c["keywords"]]
    ok = lambda w: not any(k in w for k in cat_words)  # noqa: E731
    pool = {group: [w for w in words if ok(w)] for group, words in kw["attributes"].items()}
    pool["generic"] = [w for w in kw.get("generic", []) if ok(w)]
    return pool


def _pick(rng, words: list[str]) -> str:
    return str(words[int(rng.integers(len(words)))])


def _round_price(x: float) -> int:
    """價格取到 10 元再減 1（例如 389），落在 150–990。"""
    return int(min(990, max(150, round(x / 10) * 10 - 1 if x >= 160 else 150)))


def _title(rng, noun: str, pool: dict, preorder: bool, long: bool) -> tuple[str, float]:
    n_markers = int(rng.choice([0, 1, 1, 2, 2, 3, 4]))
    markers = [str(m) for m in rng.choice(MARKERS, size=n_markers, replace=False)]
    stock_word = (_pick(rng, ["預購", "現貨+預購"]) if preorder
                  else ("現貨" if rng.random() < 0.5 else None))
    if stock_word:
        markers.append(stock_word)
    words = [_pick(rng, pool[g]) for g in ("style", "design", "material", "fit") if rng.random() < 0.7]
    if rng.random() < 0.5:
        words.append(_pick(rng, pool["generic"]))
    if long:                             # 超過 60 字：讓標題健檢有東西可以顯示
        words += [_pick(rng, pool["design"]) for _ in range(9)] + [_pick(rng, pool["style"]) for _ in range(6)]
    if rng.random() < 0.2:
        words.insert(0, _pick(rng, MARKETING))
    words.append(noun)
    front = markers[: len(markers) // 2]
    back = markers[len(markers) // 2:]
    title = " ".join(front + words + back)
    overhead = sum(len(m) for m in markers) / len(title)
    return title, overhead


def _description(rng, colors: list[str], sizes: list[str], long: bool, model: bool) -> str:
    lines = [f"顏色：{'、'.join(colors)}", f"尺寸：{'、'.join(sizes)}"]
    if long:                             # 規格段超過 150 字：觸發「描述規格段太長」的建議
        lines += [f"{s} 號：肩寬 {38 + i} 公分、胸圍 {88 + 4 * i} 公分、衣長 {60 + 2 * i} 公分、袖長 {20 + i} 公分"
                  for i, s in enumerate(["S", "M", "L", "XL"])]
    if model:
        lines.append(MODEL_INFO)
    return "\n".join(lines) + "\n\n" + NOTICE


def build_shop(seed: int = 42, kw: dict | None = None) -> list[DemoProduct]:
    """一般商品 N_REGULAR 個（依 CATEGORIES 的數量）＋ 3 個特殊商品。同一個 seed 結果完全相同。"""
    rng = np.random.default_rng(seed)
    pool = _attribute_pool(kw or config.load_yaml("keywords"))
    products: list[DemoProduct] = []
    sku_no = 0

    def add(title, category, special, preorder, overhead, base_price, n_colors, sizes,
            long_desc=False, model=False):
        nonlocal sku_no
        i = len(products) + 1
        colors = [str(c) for c in rng.choice(COLORS, size=n_colors, replace=False)]
        description = NOTICE if special else _description(rng, colors, sizes, long=long_desc, model=model)
        p = DemoProduct(pid=str(90000000000 + i * 7919), code=f"D{i:04d}", title=title, description=description,
                        category=category, special=special, preorder=preorder, overhead=overhead)
        oos_level = 0.0 if special or preorder else float(rng.choice([0, 0, 0, 0, 0, 0, 0.3, 0.3, 0.7, 1.0]))
        for color in colors:
            for size in sizes:
                sku_no += 1
                bump = 50 if size == "L" and rng.random() < 0.3 else 0
                if preorder:
                    stock = int(rng.choice(PREORDER_BASES)) - int(rng.integers(0, 9))
                else:
                    stock = 0 if rng.random() < oos_level else int(rng.integers(3, 41))
                p.skus.append(DemoSku(sku_id=str(80000000000 + sku_no * 104729), name=f"{color},{size}",
                                      sku_code=f"{p.code}-{color[0]}{size}", price=_round_price(base_price + bump),
                                      stock=stock))
        products.append(p)

    long_titles = set(int(x) for x in rng.choice(N_REGULAR, size=5, replace=False))
    long_descs = set(int(x) for x in rng.choice(N_REGULAR, size=5, replace=False))
    k = 0
    for category, (count, nouns, price) in CATEGORIES.items():
        for _ in range(count):
            preorder = rng.random() < 0.08
            title, overhead = _title(rng, _pick(rng, nouns), pool, preorder, long=k in long_titles)
            sizes = ["F"] if category == "配件" else SIZES[: int(rng.integers(1, 4))]
            n_colors = int(rng.integers(1, 3)) if len(sizes) == 3 else int(rng.integers(1, 4))
            add(title, category, False, preorder, overhead, price * rng.uniform(0.7, 1.4), n_colors, sizes,
                long_desc=k in long_descs, model=rng.random() < 0.33)
            k += 1
    for title, price in SPECIALS:
        add(title, "特殊", True, False, 0.0, price, 1, ["F"])
    return products


def _write_template(path: Path, header: list[str], rows: list[list]) -> None:
    """蝦皮批次更新範本的版面：第 0、1 列是說明，第 2 列是中文表頭，第 3–5 列是填寫說明，第 6 列起是資料。"""
    assert parsers.TEMPLATE_HEADER_ROW == 2 and parsers.TEMPLATE_DATA_ROW == 6
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ShopPilot 示範資料（虛構）"])
    ws.append([f"field_{i}" for i in range(len(header))])
    ws.append(header)
    for note in ("必填", "說明", "範例"):      # 商品ID 欄不是數字，解析時會略過
        ws.append([note] * len(header))
    for r in rows:
        ws.append(r)
    wb.save(path)


def write_catalog(path: Path, products: list[DemoProduct]) -> None:
    rows = [[p.pid, p.code, p.title, p.description] for p in products]
    _write_template(path, ["商品ID", "主商品貨號", "商品名稱", "商品描述"], rows)


def write_skus(path: Path, products: list[DemoProduct]) -> None:
    rows = []
    for p in products:
        for i, s in enumerate(p.skus):          # 主商品貨號、最低購買數量只填在每個商品的第一列（和蝦皮一樣）
            rows.append([p.pid, p.title, s.sku_id, s.name, p.code if i == 0 else None, s.sku_code, s.price, s.stock,
                         1 if i == 0 else None])
    _write_template(path, ["商品ID", "商品名稱", "商品選項ID", "商品規格名稱", "主商品貨號", "商品選項貨號", "價格",
                           "庫存", "最低購買數量"], rows)


N_NO_PERF = 10          # 本期無曝光的商品數
LOW_CTR = 4             # 曝光高、點擊率低的商品數（觸發標題相關的建議）
NO_ATC = 3              # 點擊高、加購 0 的商品數（觸發商品頁相關的建議）
N_COMPETED = 15         # 有外部競品的商品數
N_EXCLUDED = 3          # 「納入計算」填「否」的競品數


@dataclass
class DemoFiles:
    inbox: Path
    reports: list[Path]
    competitors: list[tuple[date, Path]]
    period: tuple[str, str]


def _perf_rows(rng, products: list[DemoProduct]) -> list[list]:
    """B′「最佳表現商品」：主商品列（規格ID '-'）＋規格列。數字依埋入的規律產生，再加上雜訊。"""
    regular = [p for p in products if not p.special]
    median = {c: float(np.median([p.skus[0].price for p in regular if p.category == c])) for c in CATEGORIES}
    cat_effect = {c: float(rng.normal(0, 0.25)) for c in CATEGORIES}
    hidden = set(int(x) for x in rng.choice(len(regular), size=N_NO_PERF, replace=False))
    rows: list[list] = []
    latent = []
    for i, p in enumerate(products):
        if not p.special and regular.index(p) in hidden:
            continue
        price = min(s.price for s in p.skus)
        stockout = sum(s.stock == 0 for s in p.skus) / len(p.skus)
        effect = 0.0 if p.special else cat_effect[p.category]
        imp = int(math.exp(8.3 - 3.0 * p.overhead + effect + rng.normal(0, 0.45)))
        ctr = float(np.clip(rng.normal(0.035, 0.008), 0.012, 0.07))
        ratio = 1.0 if p.special else price / median[p.category]
        conv = float(np.clip(0.07 * ratio ** -1.5 * (1 - 0.85 * stockout), 0.002, 0.4))
        latent.append([p, price, imp, ctr, conv])
    by_imp = sorted(range(len(latent)), key=lambda j: -latent[j][2])
    for j in by_imp[2: 2 + LOW_CTR]:                     # 曝光前幾名裡挑幾個，點擊率壓到很低
        latent[j][3] = 0.005
    no_atc = set(by_imp[10: 10 + NO_ATC])
    for j, (p, price, imp, ctr, conv) in enumerate(latent):
        clicks = int(rng.binomial(imp, ctr))
        visitors = int(clicks * rng.uniform(0.75, 0.9))
        orders = 0 if j in no_atc else int(rng.binomial(clicks, conv))
        units = orders + int(rng.binomial(orders, 0.15))
        atc_visitors = 0 if j in no_atc else int(rng.binomial(visitors, 0.15))
        atc_units = atc_visitors + int(rng.binomial(atc_visitors, 0.2))
        weights = rng.dirichlet(np.ones(len(p.skus)))
        sku_units = rng.multinomial(units, weights)
        sku_atc = rng.multinomial(atc_units, weights)
        sales = int(sum(u * s.price for u, s in zip(sku_units, p.skus)))
        rows.append(_perf_row(p.pid, p.title, p.code, "-", "-", sales, imp, clicks, orders, units, orders,
                              int(imp * 0.8), int(clicks * 0.9), visitors, int(visitors * 1.3),
                              int(rng.binomial(visitors, 0.42)), int(rng.binomial(clicks, 0.35)),
                              int(rng.binomial(max(visitors, 1), 0.05)), atc_visitors, atc_units))
        for s, u, a in zip(p.skus, sku_units, sku_atc):
            rows.append(_perf_row(p.pid, p.title, p.code, s.sku_id, s.name, int(u * s.price), None, None, int(u),
                                  int(u), int(u), None, None, None, None, None, None, None, int(a > 0), int(a)))
    for k, units in enumerate((2, 1)):                    # 之後被刪除的商品：不在 O 裡，沒有價格
        rows.append(_perf_row(str(99000000000 + k), DELETED_NAME, None, "-", "-", units * 290, 300, 10, units, units,
                              units, 250, 9, 8, 10, 3, 2, 0, 1, 1))
    return rows


def _perf_row(pid, name, code, sku_id, variation, sales, imp, clicks, orders, units, buyers, uimp, uclicks,
              visitors, views, bounce, search, likes, atc_visitors, atc_units) -> list:
    """依 parsers.PERF_COLS 的欄位順序排一列。"""
    values = {"商品ID": pid, "商品名稱": name, "商品當前狀態": "正常", "主商品貨號": code, "商品規格ID": sku_id,
              "商品規格": variation, "銷售額(全部訂單) (TWD)": sales, "商品曝光次數": imp, "商品點擊數": clicks,
              "全部訂單": orders, "數量(全部訂單)": units, "買家(全部訂單)": buyers, "商品卡不重複瀏覽數": uimp,
              "商品卡不重複點擊數": uclicks, "商品訪客數": visitors, "商品頁瀏覽數": views,
              "跳出商品的不重複訪客": bounce, "搜尋點擊": search, "商品按讚數": likes,
              "商品頁訪客數(加入購物車)": atc_visitors, "加入購物車(件數)": atc_units}
    return [values[c] for c in parsers.PERF_COLS]


def write_perf(path: Path, rows: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "最佳表現商品"
    ws.append(list(parsers.PERF_COLS))
    for r in rows:
        ws.append(r)
    wb.save(path)


def _competitor_name(name_rng, title: str, n: int, kw: dict) -> str:
    """從商品標題（去掉標記、行銷語後）挑 2–3 個詞當競品名稱，像真的同款商品；最後一個詞（品類）一定保留。"""
    words = core_title(title, kw).split() or title.split()
    head, noun = words[:-1], words[-1]
    k = min(len(head), int(name_rng.integers(1, 3)))
    picked = [head[int(i)] for i in sorted(name_rng.choice(len(head), size=k, replace=False))] if k else []
    return " ".join(picked + [noun]) + f" 示範競品 {n}"


def _competitor_rows(rng, products: list[DemoProduct], name_rng=None, kw: dict | None = None) -> list[dict]:
    name_rng = name_rng if name_rng is not None else np.random.default_rng(0)
    kw = kw if kw is not None else config.load_yaml("keywords")
    regular = [p for p in products if not p.special]
    chosen = [regular[int(i)] for i in sorted(rng.choice(len(regular), size=N_COMPETED, replace=False))]
    out = []
    for p in chosen:
        base = min(s.price for s in p.skus)
        for _ in range(int(rng.integers(1, 4))):
            n = len(out) + 1
            sold = int(rng.integers(50, 3000))
            out.append({"pid": p.pid, "name": _competitor_name(name_rng, p.title, n, kw),
                        "price": _round_price(base * rng.uniform(0.8, 1.25)), "sold_before": sold,
                        "sold_after": sold + int(rng.integers(5, 200)), "shop": f"示範競品賣場 {n % 7 + 1}",
                        "url": f"https://example.com/demo-competitor/{n}", "include": None})
    for i in rng.choice(len(out), size=N_EXCLUDED, replace=False):
        out[int(i)]["include"] = "否"
    return out


def write_competitors(path: Path, rows: list[dict], sold_key: str) -> None:
    """外部競品範本（ShopPilot 下載的格式）：每一列都填「我的商品ID」。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ct.SHEET
    ws.append(ct.TEMPLATE_COLUMNS)
    for r in rows:
        values = {"我的商品ID": r["pid"], "競品名稱": r["name"], "競品價格": r["price"], "已售出": r[sold_key],
                  "賣場名稱": r["shop"], "競品連結": r["url"], "納入計算": r["include"]}
        ws.append([values.get(c) for c in ct.TEMPLATE_COLUMNS])
    wb.save(path)


def generate(out_dir: Path, seed: int = 42, today: date = DEFAULT_TODAY) -> DemoFiles:
    out_dir = Path(out_dir)
    inbox, comp_dir = out_dir / "inbox", out_dir / "competitors"
    inbox.mkdir(parents=True, exist_ok=True)
    comp_dir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(seed + 1)      # 和 build_shop 分開，改商品不影響表現數字的亂數順序
    products = build_shop(seed)
    start, end = today - timedelta(days=30), today - timedelta(days=1)
    stamp = f"{today:%Y%m%d}093000"
    catalog = inbox / f"mass_update_basic_info_{SHOP_NO}_{stamp}.xlsx"
    skus = inbox / f"mass_update_sales_info_{SHOP_NO}_{stamp}.xlsx"
    perf = inbox / f"parentskudetail.{start:%Y%m%d}_{end:%Y%m%d}.xlsx"
    write_catalog(catalog, products)
    write_skus(skus, products)
    write_perf(perf, _perf_rows(rng, products))
    comp_rows = _competitor_rows(rng, products, np.random.default_rng(seed + 2))  # 名稱用另一個亂數，不影響其他數字
    competitors = []
    for day, key in ((start - timedelta(days=1), "sold_before"), (end, "sold_after")):
        path = comp_dir / ct.template_file_name(day)
        write_competitors(path, comp_rows, key)
        competitors.append((day, path))
    return DemoFiles(inbox, [catalog, skus, perf], competitors, (start.isoformat(), end.isoformat()))


def import_demo(conn, files: DemoFiles) -> list[pipeline.ImportResult]:
    """報表照一般流程匯入；競品範本用範本上的記錄日期匯入（不是今天），同期銷量才推算得出來。"""
    results = pipeline.import_folder(conn, files.inbox)
    for day, path in files.competitors:
        results.append(pipeline.import_file(conn, path.name, path.read_bytes(), today=day))
    return results


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="產生虛構賣場的蝦皮格式報表（示範資料）")
    parser.add_argument("out_dir", type=Path)
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args(argv)
    files = generate(args.out_dir, seed=args.seed)
    for p in files.reports + [p for _, p in files.competitors]:
        print(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
