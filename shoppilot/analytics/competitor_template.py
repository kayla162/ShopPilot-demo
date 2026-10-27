"""外部競品範本（半自動）：產生預先填好搜尋關鍵字的 Excel，並解析賣家填好的範本。不爬蝦皮。"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from shoppilot.analytics.products import regular_products
from shoppilot.ingest import excel_reader
from shoppilot.transform.normalize import competitor_key, parse_price_text, parse_sold_text, to_id, to_text

SHEET = "競品清單"
HELP_SHEET = "說明"
TEMPLATE_COLUMNS = ["我的商品ID", "我的商品名稱", "搜尋關鍵字", "蝦皮搜尋連結", "競品名稱", "競品價格", "已售出",
                    "賣場名稱", "競品連結", "納入計算", "上次已售出", "上次記錄日期"]
COLUMN_WIDTHS = [14, 36, 24, 44, 36, 14, 14, 16, 30, 10, 12, 14]
COMPETITOR_FIELDS = ["my_product_id", "search_keyword", "comp_name", "price_min", "price_max", "sold_count",
                     "sold_text", "shop_name", "url", "include_override", "sold_fresh"]
INCLUDE_MAP = {"是": "yes", "否": "no"}
INCLUDE_TEXT = {v: k for k, v in INCLUDE_MAP.items()}
HELP_LINES = [
    "1. 點「蝦皮搜尋連結」，在蝦皮找出和你的商品相同或相似的商品（也可以自己修改搜尋關鍵字）。",
    "2. 在該商品下方新增一列，填入「競品名稱、競品價格、已售出」；賣場名稱、競品連結可以不填。",
    "3. 同一個商品有多個競品時，「我的商品ID」留白即可，系統會沿用上一列的商品。",
    "4. 「納入計算」留白或填「是」都會納入比較；不該拿來比的競品（例如類型或價位差太多）填「否」就會排除。",
    "5. 存檔後，回到 ShopPilot 的「外部競品」分頁上傳。系統以「上傳當天」當作這次的記錄日期，請填完當天上傳。",
    "",
    "【每週更新已售出，就能和你的期間銷量比較】",
    "6. 範本會帶入目前的競品，「上次已售出」「上次記錄日期」是上次記下的數字，只供參考。",
    "7. 每週到競品頁面看一次，把目前的數字填在「已售出」；留白就沿用上次的數字，不算新的記錄。",
    "8. 系統用前後兩次「已售出」的差距，推算競品在你報表期間賣了多少件，在「競品分析」和你的銷量比較。",
    "9. 請盡量填「競品連結」：系統用連結認出同一個競品，競品改名稱也不影響；沒有連結時用賣場名稱＋競品名稱比對。",
    "注意：蝦皮把大的數字顯示成約數（例如 1.2萬）時，差距不準，這種競品不會拿來推算期間銷量。",
]


def template_file_name(today: date) -> str:
    return f"shoppilot_競品範本_{today:%Y%m%d}.xlsx"


def default_product_ids(products: pd.DataFrame, top_n: int = 20) -> list[str]:
    reg = regular_products(products)
    return reg.sort_values("impressions", ascending=False)["product_id"].head(top_n).tolist()


def _price_text(lo, hi) -> str | None:
    if lo is None or pd.isna(lo):
        return None
    fmt = lambda x: f"{x:g}"  # noqa: E731 — 299.0 → "299"
    return f"{fmt(lo)}-{fmt(hi)}" if hi is not None and pd.notna(hi) and hi > lo else fmt(lo)


def _last_sold(snapshots: pd.DataFrame | None) -> dict[tuple[str, str], tuple]:
    """(我的商品ID, 競品識別) → (最近一次記錄的已售出文字, 記錄日期)。"""
    if snapshots is None or snapshots.empty:
        return {}
    latest = snapshots.sort_values("observed_on").groupby(["my_product_id", "comp_key"]).tail(1)
    return {(r.my_product_id, r.comp_key): (r.sold_text if r.sold_text else r.sold_count, r.observed_on)
            for r in latest.itertuples()}


def build_template(products: pd.DataFrame, product_ids: list[str], competitors: pd.DataFrame | None = None,
                   snapshots: pd.DataFrame | None = None) -> bytes:
    """competitors 有值時，每個商品下方帶入目前的外部競品，「已售出」留白讓賣家填這次看到的數字。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(TEMPLATE_COLUMNS)
    index = products.set_index("product_id")
    last = _last_sold(snapshots)
    for pid in product_ids:
        row = index.loc[pid]
        code = row["code"] if isinstance(row["code"], str) else ""
        ws.append([pid, f"{code} {row['core_title']}".strip(), row["search_keyword"], row["search_url"]])
        ws.cell(row=ws.max_row, column=4).hyperlink = row["search_url"]
        if competitors is None or competitors.empty:
            continue
        for _, c in competitors[competitors["my_product_id"] == pid].iterrows():
            key = competitor_key(c.get("url"), c.get("shop_name"), c["comp_name"])
            sold, observed = last.get((pid, key), (c.get("sold_text") or c.get("sold_count"), None))
            ws.append([None, None, None, None, c["comp_name"], _price_text(c["price_min"], c.get("price_max")), None,
                       c.get("shop_name"), c.get("url"), INCLUDE_TEXT.get(c.get("include_override")),
                       None if sold is None or (isinstance(sold, float) and pd.isna(sold)) else sold, observed])
    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet(HELP_SHEET)
    for line in HELP_LINES:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_template(src) -> tuple[pd.DataFrame, list[str]]:
    try:
        rows = excel_reader.rows(src, SHEET)
    except KeyError as e:
        raise ValueError(f"找不到「{SHEET}」分頁，請使用 ShopPilot 下載的範本") from e
    if not rows:
        raise ValueError("範本是空的")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    for col in ("我的商品ID", "競品名稱"):
        if col not in header:
            raise ValueError(f"範本缺少「{col}」欄")
    index = {name: header.index(name) for name in TEMPLATE_COLUMNS if name in header}

    def cell(r: list, name: str):
        i = index.get(name)
        return r[i] if i is not None and i < len(r) else None

    records, errors = [], []
    current_id, current_kw = None, None
    for line_no, r in enumerate(rows[1:], start=2):
        pid = to_id(cell(r, "我的商品ID"))
        if pid:
            current_id, current_kw = pid, to_text(cell(r, "搜尋關鍵字"))
        name = to_text(cell(r, "競品名稱"))
        if not name:
            continue
        if not current_id:
            errors.append(f"第 {line_no} 列：找不到對應的「我的商品ID」")
            continue
        raw_price = cell(r, "競品價格")
        lo, hi = parse_price_text(raw_price)
        if lo is None:
            errors.append(f"第 {line_no} 列：無法解析價格「{raw_price}」")
            continue
        sold_raw = cell(r, "已售出")
        sold = parse_sold_text(sold_raw)
        fresh = sold is not None
        if not fresh:       # 這次沒填：沿用上次的數字（競品之間照樣能排名），但不算新的記錄
            sold_raw = cell(r, "上次已售出")
            sold = parse_sold_text(sold_raw)
        records.append({
            "my_product_id": current_id, "search_keyword": current_kw, "comp_name": name,
            "price_min": lo, "price_max": hi, "sold_count": sold, "sold_text": to_text(sold_raw),
            "shop_name": to_text(cell(r, "賣場名稱")), "url": to_text(cell(r, "競品連結")),
            "include_override": INCLUDE_MAP.get(to_text(cell(r, "納入計算")) or ""), "sold_fresh": fresh,
        })
    return pd.DataFrame(records, columns=COMPETITOR_FIELDS), errors
